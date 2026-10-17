import sys
from pptx import Presentation

prs = Presentation(r"C:\Users\Miraz\Desktop\Pocket_Drive_Presentation.pptx")
s2 = prs.slides[1]
for idx, s in enumerate(s2.shapes):
    t_list = []
    if s.has_text_frame:
        for p in s.text_frame.paragraphs:
            if p.text.strip():
                t_list.append(p.text.strip().encode('ascii', 'replace').decode('ascii'))
    print(f"Shape {idx}: name='{s.name}', left={s.left/914400:.2f}, top={s.top/914400:.2f}, w={s.width/914400:.2f}, h={s.height/914400:.2f}")
    if t_list:
        print("   Texts:", t_list)
