from pptx import Presentation

prs = Presentation(r"C:\Users\Miraz\Desktop\Pocket_Drive_Presentation.pptx")
print(f"Total Slides: {len(prs.slides)}")

s5 = prs.slides[4]
print("\n--- SLIDE 5 SHAPES ---")
for s in s5.shapes:
    t = ""
    if s.has_text_frame and s.text_frame.text.strip():
        t = s.text_frame.text.strip().replace("\n", " | ")[:60]
    print(f"  {s.name} (type={s.shape_type}): left={s.left/914400:.2f}\", top={s.top/914400:.2f}\", w={s.width/914400:.2f}\", h={s.height/914400:.2f}\" {f'-> {t}' if t else ''}")
